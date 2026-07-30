from datetime import date

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = r"D:\legal-app\Neethi_Legal_Coverage_Guide.docx"
BLUE = "153F3B"
LIGHT = "E8F1ED"
GRAY = "5B6768"
WHITE = "FFFFFF"
INK = "102B2D"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=130, bottom=100, end=130):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    total = sum(widths)
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            tc_w = cell._tc.get_or_add_tcPr().find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                cell._tc.get_or_add_tcPr().append(tc_w)
            tc_w.set(qn("w:w"), str(widths[index]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    run.font.name = "Calibri"
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor.from_string(GRAY)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    paragraph._p.append(fld)


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.78)
section.bottom_margin = Inches(0.78)
section.left_margin = Inches(1)
section.right_margin = Inches(1)
section.header_distance = Inches(0.35)
section.footer_distance = Inches(0.35)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Calibri"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
normal.font.size = Pt(10.5)
normal.font.color.rgb = RGBColor.from_string(INK)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.1

for name, size, before, after in (
    ("Heading 1", 16, 16, 8),
    ("Heading 2", 13, 12, 6),
    ("Heading 3", 11.5, 8, 4),
):
    style = styles[name]
    style.font.name = "Calibri"
    style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor.from_string(BLUE)
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.keep_with_next = True

header = section.header.paragraphs[0]
header.text = "NEETHI  |  SRI LANKAN LEGAL GUIDE"
header.style = styles["Normal"]
header.paragraph_format.space_after = Pt(0)
header_run = header.runs[0]
header_run.bold = True
header_run.font.size = Pt(8.5)
header_run.font.color.rgb = RGBColor.from_string(GRAY)
add_page_number(section.footer.paragraphs[0])

title = doc.add_paragraph()
title.paragraph_format.space_before = Pt(8)
title.paragraph_format.space_after = Pt(3)
run = title.add_run("Legal Coverage and Document Scope")
run.bold = True
run.font.name = "Calibri"
run.font.size = Pt(25)
run.font.color.rgb = RGBColor.from_string(BLUE)

subtitle = doc.add_paragraph()
subtitle.paragraph_format.space_after = Pt(12)
run = subtitle.add_run("Current capabilities of the Neethi Sri Lankan Legal AI Assistant")
run.font.size = Pt(12.5)
run.font.color.rgb = RGBColor.from_string(GRAY)

meta = doc.add_paragraph()
meta.paragraph_format.space_after = Pt(14)
run = meta.add_run(f"Prepared: {date.today().strftime('%d %B %Y')}  |  Current indexed collection")
run.bold = True
run.font.size = Pt(9)
run.font.color.rgb = RGBColor.from_string(GRAY)

callout = doc.add_table(rows=1, cols=1)
callout.style = "Table Grid"
set_table_geometry(callout, [9360])
cell = callout.cell(0, 0)
set_cell_shading(cell, LIGHT)
cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
p = cell.paragraphs[0]
p.paragraph_format.space_after = Pt(0)
r = p.add_run("Current scope")
r.bold = True
r.font.color.rgb = RGBColor.from_string(BLUE)
p.add_run(
    "  Neethi searches approximately 1,961 legal sections across 21 indexed Sri Lankan legal documents. "
    "It supports English and Sinhala questions while grounding answers in the supplied legislation."
)

doc.add_heading("1. Legal areas currently covered", level=1)

coverage = [
    ("Criminal offences", "Penal Code and Penal Code amendment Acts"),
    ("Arrests and court procedure", "Code of Criminal Procedure"),
    ("Bail and remand", "Bail Act (Consolidated 2024)"),
    ("Evidence", "Evidence Ordinance (Consolidated 2024)"),
    ("Theft, robbery and physical injury", "Penal Code"),
    ("Domestic violence", "Prevention of Domestic Violence Act No. 34 of 2005"),
    ("Road and vehicle offences", "Motor Traffic Act"),
    ("Computer crimes", "Computer Crime Act No. 24 of 2007"),
    ("Online offences", "Online Safety Act No. 9 of 2024"),
    ("Data privacy", "Personal Data Protection Act No. 9 of 2022"),
    ("Consumer matters", "Consumer Affairs Authority Act No. 9 of 2003"),
    ("Employment disputes", "Industrial Disputes Act (Consolidated 2024)"),
    ("Working hours and conditions", "Shop and Office Employees Act (Consolidated 2024)"),
    ("Wages", "Wages Boards Ordinance and Minimum Wages (Indian Labour) Ordinance"),
    ("Electronic records and signatures", "Electronic Transactions Act and 2017 amendment"),
    ("Victims and witnesses", "Assistance to and Protection of Victims of Crime and Witnesses Act No. 10 of 2023"),
]

table = doc.add_table(rows=1, cols=2)
table.style = "Table Grid"
table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
headers = table.rows[0].cells
headers[0].text = "Legal area"
headers[1].text = "Documents currently included"
for cell in headers:
    set_cell_shading(cell, BLUE)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for run in cell.paragraphs[0].runs:
        run.bold = True
        run.font.color.rgb = RGBColor.from_string(WHITE)
        run.font.size = Pt(9.5)
for area, documents in coverage:
    cells = table.add_row().cells
    cells[0].text = area
    cells[1].text = documents
    for cell in cells:
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        cell.paragraphs[0].paragraph_format.space_after = Pt(0)
        for run in cell.paragraphs[0].runs:
            run.font.size = Pt(9)
    cells[0].paragraphs[0].runs[0].bold = True
set_table_geometry(table, [2800, 6560])

doc.add_heading("2. Current limitations", level=1)
limitations = [
    "The system mainly searches legislation; it does not yet provide comprehensive court judgments or case law.",
    "The collection does not contain every Sri Lankan law or every amendment.",
    "Sinhala questions are supported, but most underlying legal documents are currently in English.",
    "Family law matters such as divorce, maintenance and child custody are not yet covered in detail.",
    "Land law, inheritance, taxation, company law, immigration and constitutional law need broader source collections.",
    "Neethi provides general legal information. It does not give a final legal opinion, decide whether an offence occurred, or predict a court outcome.",
]
for item in limitations:
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    p.paragraph_format.space_after = Pt(6)
    p.add_run(item)

doc.add_heading("3. Recommended next document priorities", level=1)
priorities = [
    ("Marriage, divorce, maintenance and child custody", "Expand practical family-law coverage."),
    ("Land, property and tenancy", "Support ownership, lease, possession and landlord-tenant questions."),
    ("Inheritance and wills", "Cover succession, probate and estate-related questions."),
    ("Fundamental rights and constitutional law", "Support rights-based and public-law questions."),
    ("Labour tribunal decisions", "Add practical interpretation to the existing employment legislation."),
    ("Supreme Court and Court of Appeal judgments", "Introduce authoritative case-law context and precedent."),
    ("Sinhala versions of existing legislation", "Improve Sinhala terminology and allow source text to be presented directly in Sinhala."),
]
for index, (title_text, description) in enumerate(priorities, start=1):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(title_text)
    r.bold = True
    p.add_run(f" - {description}")

doc.add_heading("4. Document quality requirements", level=1)
doc.add_paragraph(
    "For reliable retrieval, new documents should preferably be official or authoritative PDF copies with clear section numbering. "
    "Each source should be recorded with its publication date, effective date, amendment status and official source URL where available."
)

note = doc.add_table(rows=1, cols=1)
note.style = "Table Grid"
set_table_geometry(note, [9360])
cell = note.cell(0, 0)
set_cell_shading(cell, "F5F1E8")
p = cell.paragraphs[0]
p.paragraph_format.space_after = Pt(0)
r = p.add_run("Important: ")
r.bold = True
r.font.color.rgb = RGBColor.from_string(BLUE)
p.add_run(
    "Coverage indicates that relevant source material is available for retrieval. It does not guarantee that the system can answer every factual scenario."
)

doc.core_properties.title = "Neethi Legal Coverage and Document Scope"
doc.core_properties.subject = "Current Sri Lankan legal areas, indexed documents, limitations and expansion priorities"
doc.core_properties.author = "Neethi Project"
doc.save(OUTPUT)
print(OUTPUT)
